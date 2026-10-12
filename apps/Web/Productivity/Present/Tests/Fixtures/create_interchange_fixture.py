"""Independent OOXML fixture, authored with python-pptx 1.0.2; no 9to1 exporter involved."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

out = Path(__file__).with_name('external-styled-notes.pptx')
r = Presentation()
r.slide_width, r.slide_height = Inches(13.333333), Inches(7.5)
s = r.slides.add_slide(r.slide_layouts[5])
s.shapes.title.text = 'Quarterly roadmap'
box = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(2.4), Inches(2.25), Inches(6.4), Inches(1.75))
box.fill.solid()
box.fill.fore_color.rgb = RGBColor(68, 114, 196)
box.line.color.rgb = RGBColor(32, 64, 128)
box.line.width = Pt(2)
text = s.shapes.add_textbox(Inches(2.6), Inches(2.4), Inches(6), Inches(1.45))
f = text.text_frame
f.word_wrap = True
f.vertical_anchor = MSO_ANCHOR.MIDDLE
f.margin_left = f.margin_right = f.margin_top = f.margin_bottom = 0
for i, value in enumerate(['Plan together.', 'Build clearly.']):
    p = f.paragraphs[0] if i == 0 else f.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = value
    run.font.name = 'Montserrat'
    run.font.size = Pt(36)
    run.font.bold = True
    run.font.italic = True
    run.font.underline = True
    run.font.color.rgb = RGBColor(255, 255, 255)
s.notes_slide.notes_text_frame.text = 'Present the priorities.\nPause for questions.\nUnicode: café 👩🏽‍💻'
second = r.slides.add_slide(r.slide_layouts[5])
second.shapes.title.text = 'Appendix'
second._element.set('show', '0')
r.save(out)
print(out)
